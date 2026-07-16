# UniBus Live — Thesis Defense deck (16:9, minimal, image-forward)
# All facts from the real pilot. Fonts: Segoe UI (universal on Windows).
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image
import os

IMG = "img"
EMU_IN = 914400

# ---------- palette ----------
INK    = RGBColor(0x1A, 0x25, 0x38)
BODY   = RGBColor(0x5A, 0x6B, 0x85)
MUTED  = RGBColor(0x8A, 0x97, 0xAB)
NAVY   = RGBColor(0x1F, 0x3A, 0x6E)
BLUE   = RGBColor(0x25, 0x63, 0xEB)
TEAL   = RGBColor(0x0F, 0x76, 0x6E)
AMBER  = RGBColor(0xD9, 0x77, 0x06)
GREEN  = RGBColor(0x15, 0x80, 0x3D)
VIOLET = RGBColor(0x7C, 0x3A, 0xED)
SLATE  = RGBColor(0x47, 0x55, 0x69)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
PANEL  = RGBColor(0xF2, 0xF6, 0xFB)   # soft blue-gray card
PANEL2 = RGBColor(0xFB, 0xFC, 0xFE)
LINEC  = RGBColor(0xDD, 0xE5, 0xF0)
AMBER_BG = RGBColor(0xFD, 0xF3, 0xE4)
GREEN_BG = RGBColor(0xEA, 0xF6, 0xEE)
BLUE_BG  = RGBColor(0xEB, 0xF1, 0xFD)

FONT = "Segoe UI"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def slide():
    return prs.slides.add_slide(BLANK)

def rect(s, x, y, w, h, fill=None, line=None, line_w=0.75, radius=0.08, shadow=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp

def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=4, line_spacing=1.0, wrap=True):
    """runs: list of paragraphs; each paragraph = list of (txt, size, color, bold)"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        for (t, size, color, bold) in para:
            r = p.add_run(); r.text = t
            r.font.name = FONT; r.font.size = Pt(size)
            r.font.color.rgb = color; r.font.bold = bold
    return tb

def pic(s, path, x, y, w=None, h=None, border=True):
    im = Image.open(os.path.join(IMG, path))
    ar = im.width / im.height
    if w is None: w = h * ar
    if h is None: h = w / ar
    p = s.shapes.add_picture(os.path.join(IMG, path), Inches(x), Inches(y),
                             Inches(w), Inches(h))
    if border:
        p.line.color.rgb = LINEC; p.line.width = Pt(1)
    p.shadow.inherit = False
    return (w, h)

def header(s, kicker, title, accent, title_size=31):
    rect(s, 0.75, 0.62, 0.28, 0.28, fill=accent, radius=0.5)
    text(s, 1.18, 0.60, 10.5, 0.4,
         [[(kicker.upper(), 12.5, accent, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.75, 0.98, 11.9, 0.9,
         [[(title, title_size, INK, True)]])

def footer(s, n):
    text(s, 0.75, 7.08, 6, 0.3,
         [[("UniBus Live · Thesis Defense", 9, MUTED, False)]])
    text(s, 12.15, 7.08, 0.45, 0.3,
         [[(str(n), 9, MUTED, False)]], align=PP_ALIGN.RIGHT)

def stat_tile(s, x, y, w, h, value, label, vcolor=NAVY, bg=PANEL, vsize=30):
    rect(s, x, y, w, h, fill=bg)
    text(s, x, y + 0.13, w, 0.62, [[(value, vsize, vcolor, True)]],
         align=PP_ALIGN.CENTER)
    text(s, x + 0.1, y + h - 0.62, w - 0.2, 0.55,
         [[(label, 10.5, BODY, False)]], align=PP_ALIGN.CENTER, line_spacing=0.95)

# ============================================================ 1 · TITLE
s = slide()
rect(s, 0, 0, 13.333, 7.5, fill=WHITE, radius=0)
rect(s, 0, 0, 13.333, 0.14, fill=NAVY, radius=0)
pic(s, "logo.png", 6.06, 0.62, w=1.22, border=False)
text(s, 0.9, 2.02, 11.53, 0.4,
     [[("B.SC. THESIS DEFENSE  ·  DEPARTMENT OF CSE  ·  NPI UNIVERSITY OF BANGLADESH",
        12.5, BLUE, True)]], align=PP_ALIGN.CENTER)
text(s, 0.9, 2.5, 11.53, 1.6,
     [[("UniBus Live", 46, NAVY, True)],
      [("A Real-Time Bilingual University Bus Tracking System", 21, INK, False)],
      [("with Reliable Dual-Source GPS Acquisition", 21, INK, False)]],
     align=PP_ALIGN.CENTER, space_after=6)
# team cards
names = [("Md. Shishir Hossain", "ID 0852220005101005"),
         ("Md. Sajib Hasan",     "ID 0852220005101008"),
         ("Md. Showkat Hossen",  "ID 0852220005101009")]
tw = 3.3; gap = 0.3; x0 = (13.333 - (3*tw + 2*gap)) / 2
for i, (nm, idn) in enumerate(names):
    x = x0 + i*(tw+gap)
    rect(s, x, 4.62, tw, 0.92, fill=PANEL)
    text(s, x, 4.76, tw, 0.4, [[(nm, 14, INK, True)]], align=PP_ALIGN.CENTER)
    text(s, x, 5.12, tw, 0.32, [[(idn + "  ·  Batch 12th", 10.5, BODY, False)]],
         align=PP_ALIGN.CENTER)
text(s, 0.9, 5.95, 11.53, 0.75,
     [[("Supervisor: Mahbubur Rahman Chowdhury, Senior Lecturer, Dept. of CSE, NPIUB",
        13, BODY, False)],
      [("18 July 2026", 13, BLUE, True)]],
     align=PP_ALIGN.CENTER, space_after=5)

# ============================================================ 2 · PROBLEM
s = slide()
header(s, "The Problem", "Waiting for a bus you cannot see", AMBER)
text(s, 0.75, 1.95, 11.8, 0.75,
     [[("A student at the stop has no way to know: did the bus leave 5 minutes ago — "
        "or is it 15 minutes away?", 17, BODY, False)]], line_spacing=1.1)
cards = [("Missed buses", "The bus passed early; the student finds out too late.", "⏱"),
         ("Long, uncertain waits", "Standing in heat or rain for time that could be saved.", "🌧"),
         ("No usable product", "Commercial trackers: costly, operator-facing, English-only.", "💸")]
cw = 3.83; x0 = 0.75
for i, (t1, t2, ic) in enumerate(cards):
    x = x0 + i*(cw+0.155)
    rect(s, x, 2.95, cw, 2.5, fill=AMBER_BG)
    text(s, x+0.3, 3.25, cw-0.6, 0.5, [[(ic, 22, AMBER, False)]])
    text(s, x+0.3, 3.85, cw-0.6, 0.5, [[(t1, 17, INK, True)]])
    text(s, x+0.3, 4.38, cw-0.6, 0.95, [[(t2, 12.5, BODY, False)]], line_spacing=1.1)
text(s, 0.75, 5.85, 11.8, 0.6,
     [[("Our goal: ", 15, INK, True),
       ("let every student see their bus live, with the arrival time at their own stop — in their own language.",
        15, BODY, False)]], line_spacing=1.1)
footer(s, 2)

# ============================================================ 3 · SOLUTION
s = slide()
header(s, "Our Solution", "UniBus Live — one platform, four parts", BLUE)
comp = [("GPS Tracker", "Concox GT06S fixed in the bus.\nFully automatic — no driver needed.", "PRIMARY SOURCE", GREEN),
        ("Passenger App", "Live map · per-stop arrival times ·\nalerts · English + বাংলা", "FOR STUDENTS", BLUE),
        ("Driver App", "Opt-in backup publisher —\nused only when needed.", "BACKUP SOURCE", AMBER),
        ("Admin Console", "Routes · schedules · fleet ·\nlive operations map.", "FOR TRANSPORT OFFICE", VIOLET)]
cw = 2.83; x0 = 0.75
for i, (t1, t2, tag, ac) in enumerate(comp):
    x = x0 + i*(cw+0.17)
    rect(s, x, 2.1, cw, 3.15, fill=PANEL)
    rect(s, x, 2.1, cw, 0.12, fill=ac, radius=0)
    text(s, x+0.25, 2.42, cw-0.5, 0.35, [[(tag, 9.5, ac, True)]])
    text(s, x+0.25, 2.8, cw-0.5, 0.55, [[(t1, 18, INK, True)]])
    for j, ln in enumerate(t2.split("\n")):
        text(s, x+0.25, 3.42+j*0.62, cw-0.5, 0.6, [[(ln, 11.5, BODY, False)]],
             line_spacing=1.05)
rect(s, 0.75, 5.62, 11.83, 0.85, fill=BLUE_BG)
text(s, 0.75, 5.62, 11.83, 0.85,
     [[("Everything runs on ", 14.5, INK, False),
       ("one low-cost server", 14.5, NAVY, True),
       (" — Node.js · PostgreSQL · Redis · Traccar · Socket.IO. No paid cloud services.",
        14.5, INK, False)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 3)

# ============================================================ 4 · PASSENGER APP
s = slide()
header(s, "The Passenger Experience", "The map is never blank", BLUE)
shots = ["p1.jpg", "p2.jpg", "p3.jpg", "p4.jpg"]
ph = 4.15; pw = ph * (738/1600)
x0 = 0.75
for i, sh in enumerate(shots):
    pic(s, sh, x0 + i*(pw+0.22), 2.0, h=ph)
capx = x0 + 4*(pw+0.22) + 0.15
caps = [("Before departure", "Bus visible at the depot — \"getting ready\", not a blank map."),
        ("On route, live", "Next stop and minutes away, updated in seconds."),
        ("Today's schedules", "Every route and departure at a glance."),
        ("English + বাংলা", "The whole app in the rider's language.")]
for i, (t1, t2) in enumerate(caps):
    y = 2.0 + i*1.06
    rect(s, capx, y, 0.09, 0.9, fill=BLUE, radius=0.5)
    text(s, capx+0.25, y, 12.58-capx-0.25, 0.35, [[(t1, 13.5, INK, True)]])
    text(s, capx+0.25, y+0.34, 12.58-capx-0.25, 0.6, [[(t2, 10.5, BODY, False)]],
         line_spacing=1.0)
text(s, 0.75, 6.35, 11.8, 0.4,
     [[("Get-off alert: a push notification when the bus is within 80 m of your stop — even if the app is closed.",
        13, BODY, False)]])
footer(s, 4)

# ============================================================ 5 · ARCHITECTURE
s = slide()
header(s, "System Architecture", "Three layers, one low-cost server", TEAL)
pic(s, "arch.png", 0.75, 1.95, h=4.9)
cx = 7.3
pts = [("Vehicle layer", "Tracker reports every 5 s over GSM; driver phone is only a fallback."),
       ("One backend server", "Traccar decodes · Node.js validates, computes ETAs, broadcasts · PostgreSQL + Redis."),
       ("Client layer", "Passenger app and admin console — pushed over WebSocket, not polled.")]
for i, (t1, t2) in enumerate(pts):
    y = 2.35 + i*1.5
    rect(s, cx, y, 5.28, 1.28, fill=PANEL)
    text(s, cx+0.28, y+0.17, 4.8, 0.4, [[(t1, 15, TEAL, True)]])
    text(s, cx+0.28, y+0.56, 4.8, 0.68, [[(t2, 11.5, BODY, False)]], line_spacing=1.05)
footer(s, 5)

# ============================================================ 6 · DATA FLOW
s = slide()
header(s, "Real-Time Data Flow", "From GPS fix to the student's screen", TEAL)
pic(s, "flow.png", 0.75, 1.9, h=5.0)
cx = 6.7
rect(s, cx, 2.5, 5.85, 1.7, fill=GREEN_BG)
text(s, cx+0.35, 2.78, 5.2, 0.6, [[("5.9 s", 40, GREEN, True),
                                    ("  median", 16, BODY, False)]])
text(s, cx+0.35, 3.62, 5.2, 0.45,
     [[("from GPS fix on the bus to the server — measured over 46,286 real fixes.",
        12.5, BODY, False)]], line_spacing=1.05)
steps = [("1–2", "Capture & decode — every 5 s, decoded in ≈1 s"),
         ("3–4", "Pick best source · reject faulty readings"),
         ("5",   "Compute the ETA for every stop ahead"),
         ("6–7", "Store, publish, push to every subscribed phone")]
for i, (n, t) in enumerate(steps):
    y = 4.55 + i*0.55
    text(s, cx, y, 0.75, 0.4, [[(n, 13, TEAL, True)]])
    text(s, cx+0.8, y, 5.1, 0.4, [[(t, 12.5, BODY, False)]])
footer(s, 6)

# ============================================================ 7 · DUAL SOURCE
s = slide()
header(s, "Reliable by Design", "Two independent sources, one automatic choice", VIOLET)
rect(s, 0.75, 2.0, 5.6, 3.1, fill=GREEN_BG)
text(s, 1.05, 2.28, 5.0, 0.4, [[("PRIMARY — GPS TRACKER", 12, GREEN, True)]])
text(s, 1.05, 2.7, 5.0, 0.5, [[("Works completely on its own", 17.5, INK, True)]])
for i, t in enumerate(["Hard-wired in the bus, reports every 5 s",
                       "Trip starts automatically from its data",
                       "No driver action, phone, or data plan needed"]):
    text(s, 1.05, 3.3+i*0.5, 5.1, 0.45, [[("•  ", 12.5, GREEN, True), (t, 12.5, BODY, False)]])
rect(s, 6.98, 2.0, 5.6, 3.1, fill=AMBER_BG)
text(s, 7.28, 2.28, 5.0, 0.4, [[("BACKUP — DRIVER'S PHONE (OPT-IN)", 12, AMBER, True)]])
text(s, 7.28, 2.7, 5.0, 0.5, [[("Only when a trip needs it", 17.5, INK, True)]])
for i, t in enumerate(["Driver must be assigned and start sharing",
                       "Covers buses without a tracker yet",
                       "Normal operation never depends on it"]):
    text(s, 7.28, 3.3+i*0.5, 5.1, 0.45, [[("•  ", 12.5, AMBER, True), (t, 12.5, BODY, False)]])
rect(s, 0.75, 5.42, 11.83, 1.05, fill=PANEL)
text(s, 0.75, 5.42, 11.83, 1.05,
     [[("The server scores both sources on every fix and picks the best one automatically",
        15.5, INK, True)],
      [("tracker preferred when healthy · switches mid-trip with no visible jump · what is automatic is the selection",
        12, BODY, False)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=3)
footer(s, 7)

# ============================================================ 8 · ETA
s = slide()
header(s, "Smart Arrival Times", "A simple estimator that works from day one", VIOLET)
rect(s, 0.75, 2.1, 11.83, 1.5, fill=PANEL)
text(s, 0.75, 2.1, 11.83, 1.5,
     [[("ETA  =  remaining distance along the route  ÷  recent average speed",
        22, NAVY, True)],
      [("computed for every stop ahead, on every accepted GPS fix", 13, BODY, False)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=6)
feats = [("Route-aware", "The fix is projected onto the route line, so sideways GPS error doesn't break the estimate."),
         ("Honest confidence", "Every estimate is labeled High / Medium / Low from the quality of the speed signal."),
         ("No training data", "Works from the first day — while logging every prediction so an ML model can come later.")]
cw = 3.83
for i, (t1, t2) in enumerate(feats):
    x = 0.75 + i*(cw+0.155)
    rect(s, x, 3.95, cw, 1.85, fill=PANEL2, line=LINEC)
    text(s, x+0.28, 4.2, cw-0.56, 0.4, [[(t1, 15, VIOLET, True)]])
    text(s, x+0.28, 4.66, cw-0.56, 1.0, [[(t2, 11.5, BODY, False)]], line_spacing=1.1)
text(s, 0.75, 6.15, 11.8, 0.45,
     [[("Arrival radius 80 m — reaching it marks the stop and fires the rider's get-off alert.",
        13.5, BODY, False)]], align=PP_ALIGN.CENTER)
footer(s, 8)

# ============================================================ 9 · PILOT
s = slide()
header(s, "The Pilot", "Seven weeks on the real fleet — zero simulation", GREEN)
tiles = [("7", "weeks live\n6 May – 23 Jun 2026"), ("9", "routes"),
         ("3", "buses (1 with\nhardware tracker)"), ("38", "registered riders"),
         ("71", "completed trips"), ("46,548", "GPS reports ingested")]
tw2 = 3.83; th = 1.95
for i, (v, l) in enumerate(tiles):
    x = 0.75 + (i % 3)*(tw2+0.155)
    y = 2.15 + (i // 3)*(th+0.25)
    rect(s, x, y, tw2, th, fill=PANEL)
    text(s, x, y+0.28, tw2, 0.8, [[(v, 36, NAVY, True)]], align=PP_ALIGN.CENTER)
    text(s, x+0.2, y+1.12, tw2-0.4, 0.7,
         [[(ln, 11.5, BODY, False)] for ln in l.split("\n")],
         align=PP_ALIGN.CENTER, space_after=1, line_spacing=1.0)
text(s, 0.75, 6.55, 11.8, 0.4,
     [[("Every number in this defense comes from the production database of this pilot.",
        13, BODY, False)]], align=PP_ALIGN.CENTER)
footer(s, 9)

# ============================================================ 10 · SPEED
s = slide()
header(s, "Result · Speed", "The map reflects reality within seconds", GREEN)
text(s, 0.75, 2.05, 5.3, 1.4, [[("5.9 s", 64, BLUE, True)]])
text(s, 0.75, 3.35, 5.3, 0.5,
     [[("median latency, GPS fix → server", 14.5, BODY, False)]])
text(s, 0.75, 3.95, 5.4, 1.3,
     [[("Bounded mainly by the 8 s polling interval — push delivery itself takes ≈1 s. "
        "A further 1.2 s (median) to persist each fix.", 12.5, BODY, False)]],
     line_spacing=1.15)
# percentile bars — one measure, one hue
bx, bw_max, bh = 7.05, 4.4, 0.52
vals = [("Median", 5.9), ("Mean", 7.4), ("90th percentile", 9.9), ("95th percentile", 11.1)]
vmax = 11.1
text(s, bx, 2.0, 5.3, 0.4, [[("Latency across 46,286 fixes (seconds)", 13, INK, True)]])
for i, (lab, v) in enumerate(vals):
    y = 2.6 + i*(bh+0.42)
    text(s, bx, y-0.02, 5.3, 0.3, [[(lab, 11.5, BODY, False)]])
    w = bw_max * (v/vmax)
    rect(s, bx, y+0.26, w, bh-0.24, fill=BLUE, radius=0.35)
    text(s, bx + w + 0.12, y+0.22, 1.0, 0.35, [[(f"{v} s", 12.5, INK, True)]])
footer(s, 10)

# ============================================================ 11 · RELIABILITY
s = slide()
header(s, "Result · Reliability & Engagement", "The backup carried real weight", GREEN)
# split bar 64/36
text(s, 0.75, 2.1, 11.8, 0.4, [[("Which source tracked the 70 sourced trips?", 14.5, INK, True)]])
bx, by, bw_full, bh2 = 0.75, 2.62, 11.83, 0.85
w1 = bw_full * 0.64
rect(s, bx, by, w1 - 0.02, bh2, fill=BLUE, radius=0.12)
rect(s, bx + w1 + 0.02, by, bw_full - w1 - 0.02, bh2, fill=AMBER, radius=0.12)
text(s, bx+0.35, by+0.16, 4.5, 0.55,
     [[("64%  GPS tracker  (45 trips)", 15, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
text(s, bx+w1+0.3, by+0.16, 4.0, 0.55,
     [[("36%  phone backup  (25)", 15, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.75, 3.62, 11.8, 0.45,
     [[("The 36% were mostly the two buses not yet fitted with trackers — coverage never stopped. "
        "Fleet-wide trackers will shrink this to a rare safety net.", 12.5, BODY, False)]],
     line_spacing=1.1)
stat_tile(s, 0.75, 4.55, 3.83, 1.85, "80%", "of arrival & service notifications\nwere read by riders (36 of 45)", GREEN)
stat_tile(s, 4.74, 4.55, 3.83, 1.85, "80 m", "arrival radius that triggers\nthe get-off alert", NAVY)
stat_tile(s, 8.73, 4.55, 3.85, 1.85, "0", "passenger-visible jumps during\nmid-trip source handovers", TEAL)
footer(s, 11)

# ============================================================ 12 · HIDDEN PROBLEM
s = slide()
header(s, "A Hidden Problem — Found & Fixed", "The pilot's most valuable discovery", AMBER)
rect(s, 0.75, 2.05, 5.6, 3.3, fill=AMBER_BG)
text(s, 0.75, 2.35, 5.6, 1.0, [[("98%", 58, AMBER, True)]], align=PP_ALIGN.CENTER)
text(s, 1.1, 3.6, 4.9, 1.5,
     [[("of 46,548 GPS reports were pure waste", 15, INK, True)],
      [("buses kept transmitting every 5 seconds even while parked overnight — "
        "only 864 reports fell inside a real trip", 12, BODY, False)]],
     align=PP_ALIGN.CENTER, space_after=6, line_spacing=1.1)
text(s, 6.35, 3.3, 0.65, 0.7, [[("→", 34, SLATE, True)]], align=PP_ALIGN.CENTER)
rect(s, 6.98, 2.05, 5.6, 3.3, fill=GREEN_BG)
text(s, 6.98, 2.35, 5.6, 1.0, [[("60×", 58, GREEN, True)]], align=PP_ALIGN.CENTER)
text(s, 7.33, 3.6, 4.9, 1.5,
     [[("less parked-time reporting after our fix", 15, INK, True)],
      [("device: 5 s → 300 s when parked (one SMS: TIMER,5,300#)  ·  "
        "server polling: 8 s → 60 s when idle", 12, BODY, False)]],
     align=PP_ALIGN.CENTER, space_after=6, line_spacing=1.1)
text(s, 0.75, 5.75, 11.8, 0.85,
     [[("Invisible in design, unmistakable in real data — and fixed at both the device and the server ",
        13.5, BODY, False),
       ("without touching live tracking.", 13.5, INK, True)]],
     align=PP_ALIGN.CENTER, line_spacing=1.1)
footer(s, 12)

# ============================================================ 13 · ADMIN
s = slide()
header(s, "The Admin Console", "The transport office sees everything", VIOLET)
pic(s, "a1.jpg", 0.75, 2.0, w=5.75)
pic(s, "a2.jpg", 6.82, 2.0, w=5.75)
text(s, 0.75, 5.45, 5.75, 0.4, [[("Live operations — whole fleet at a glance", 12.5, BODY, False)]],
     align=PP_ALIGN.CENTER)
text(s, 6.82, 5.45, 5.75, 0.4, [[("Route builder — stops, distance, travel time on the map", 12.5, BODY, False)]],
     align=PP_ALIGN.CENTER)
text(s, 0.75, 6.15, 11.8, 0.5,
     [[("Also: buses, GPS devices, schedules, stops, users, analytics — one web console, same API as the apps.",
        13, BODY, False)]], align=PP_ALIGN.CENTER)
footer(s, 13)

# ============================================================ 14 · DRIVER APP
s = slide()
header(s, "The Driver App", "The opt-in backup, kept deliberately simple", AMBER)
dh = 4.3; dw = dh * (738/1600)
x0 = 0.75
for i, sh in enumerate(["d1.jpg", "d2.jpg", "d3.jpg"]):
    pic(s, sh, x0 + i*(dw+0.3), 2.0, h=dh)
cx = x0 + 3*(dw+0.3) + 0.25
pts = [("Shares nothing by default", "Location goes out only between Start trip and End trip."),
       ("Choose the source", "Bus device or my phone — visible on screen."),
       ("Live self-check", "Driver sees speed, GPS accuracy and update age while publishing."),
       ("Drivers only", "Assigned driver + driver role required — an authorization boundary.")]
for i, (t1, t2) in enumerate(pts):
    y = 2.0 + i*1.12
    rect(s, cx, y, 0.09, 0.95, fill=AMBER, radius=0.5)
    text(s, cx+0.25, y, 12.58-cx-0.25, 0.35, [[(t1, 13.5, INK, True)]])
    text(s, cx+0.25, y+0.35, 12.58-cx-0.25, 0.7, [[(t2, 10.5, BODY, False)]], line_spacing=1.0)
footer(s, 14)

# ============================================================ 15 · LIMITS & FUTURE
s = slide()
header(s, "Honest Limits & What's Next", "A feasibility pilot, and a roadmap", SLATE)
rect(s, 0.75, 2.0, 5.6, 4.45, fill=PANEL)
text(s, 1.05, 2.25, 5.0, 0.4, [[("LIMITATIONS", 12.5, SLATE, True)]])
lims = ["Small pilot: 7 weeks, 3 buses, 38 riders",
        "Only one bus carried the hardware tracker",
        "ETA error (MAE) still accumulating — every prediction is logged for it",
        "Timetable linkage incomplete during the pilot"]
for i, t in enumerate(lims):
    text(s, 1.05, 2.75+i*0.85, 5.05, 0.8,
         [[("•  ", 12.5, SLATE, True), (t, 12.5, BODY, False)]], line_spacing=1.1)
rect(s, 6.98, 2.0, 5.6, 4.45, fill=BLUE_BG)
text(s, 7.28, 2.25, 5.0, 0.4, [[("FUTURE WORK", 12.5, BLUE, True)]])
futs = ["ML arrival prediction trained on our logged trip history",
        "Trackers on every bus — phone becomes a rare safety net",
        "Ignition-wired (ACC) reporting: device sleeps with the engine",
        "Occupancy info and a formal usability study at larger scale"]
for i, t in enumerate(futs):
    text(s, 7.28, 2.75+i*0.85, 5.05, 0.8,
         [[("•  ", 12.5, BLUE, True), (t, 12.5, BODY, False)]], line_spacing=1.1)
footer(s, 15)

# ============================================================ 16 · THANKS
s = slide()
rect(s, 0, 0, 13.333, 7.5, fill=NAVY, radius=0)
text(s, 0.9, 2.2, 11.53, 1.0, [[("Thank you", 52, WHITE, True)]], align=PP_ALIGN.CENTER)
text(s, 0.9, 3.35, 11.53, 0.5,
     [[("Questions welcome — every number is from our production database.",
        16, RGBColor(0xC9, 0xD6, 0xEC), False)]], align=PP_ALIGN.CENTER)
strip = [("5.9 s", "median latency"), ("46,548", "GPS reports"), ("71", "real trips"),
         ("64 / 36", "tracker vs phone"), ("60×", "idle waste cut"), ("80%", "alerts read")]
tw3 = 1.85; x0 = (13.333 - 6*tw3 - 5*0.12) / 2
for i, (v, l) in enumerate(strip):
    x = x0 + i*(tw3+0.12)
    rect(s, x, 4.35, tw3, 1.25, fill=RGBColor(0x2A, 0x4A, 0x86))
    text(s, x, 4.52, tw3, 0.5, [[(v, 20, WHITE, True)]], align=PP_ALIGN.CENTER)
    text(s, x, 5.08, tw3, 0.4, [[(l, 9.5, RGBColor(0xC9, 0xD6, 0xEC), False)]],
         align=PP_ALIGN.CENTER)
text(s, 0.9, 6.1, 11.53, 0.5,
     [[("UniBus Live · Team UniBus · Dept. of CSE · NPI University of Bangladesh",
        12, RGBColor(0x9A, 0xAC, 0xCC), False)]], align=PP_ALIGN.CENTER)

prs.save("UniBus_Live_Defense.pptx")
print("saved")
print("slides:", len(prs.slides._sldIdLst))
